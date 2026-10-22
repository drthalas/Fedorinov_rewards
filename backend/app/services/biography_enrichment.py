"""Biography-only enrichment. No research, schema changes, inserts or media writes.

Signed export snapshots and private preview/report files live outside SQLite in
the OS temporary directory. Clearing that directory invalidates old exports.
"""
from collections import Counter
from contextlib import closing
from io import BytesIO
from pathlib import Path
from urllib.parse import urlsplit
from zipfile import ZipFile
import hashlib
import hmac
import json
import os
import re
import secrets
import tempfile
import time

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from ..db import open_readonly_connection, open_write_connection
from .write_guard import ensure_write_allowed

MAX_BYTES = 10 * 1024 * 1024
MAX_UNPACKED = 50 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLUMNS = 40
TTL = 7 * 86400
IDS = ('id_gos', 'id_catigory', 'id_sub_catigory', 'id_name')
HEADERS = ['person_id', 'ФИО', 'Год рождения', 'Звание', 'Награда', 'Иерархия ID',
           'Номер(а) награды', 'Примечание к году', 'biography', 'source_url', 'status', 'notes',
           'ФИО в карточке', 'Год в карточке', 'Награда в карточке', 'Звание в карточке',
           'Идентификация', 'snapshot', 'snapshot_signature']
VISIBLE_HEADERS = ['ID кавалера', 'ФИО', 'Год рождения', 'Награда', 'Биография', 'Источник', 'Статус']
METADATA_SHEET = 'Служебные данные'
ALIASES = {'ID кавалера': 'person_id', 'Источник': 'source_url',
           'Биография': 'biography', 'WarHeroes URL': 'source_url', 'Статус': 'status',
           'Подтверждённые факты': 'notes'}
READY = 'Готово'
SKIP = 'Пропуск'
CONFLICT = 'Конфликт'
ERROR = 'Ошибка'
READY_LABEL = 'Готово к сохранению'
REVIEW_LABEL = 'Требует проверки'
SKIP_LABEL = 'Пропущено'


def review_groups(record):
    groups = {name: [] for name in (READY_LABEL, REVIEW_LABEL, SKIP_LABEL)}
    for row in record['rows']:
        if row['kind'] == 'Сохранено':
            continue
        label = READY_LABEL if row['kind'] == READY else SKIP_LABEL if row['kind'] == SKIP else REVIEW_LABEL
        groups[label].append(row)
    return groups


def category(row):
    if row['kind'] == 'Сохранено':
        return 'Сохранено'
    return READY_LABEL if row['kind'] == READY else SKIP_LABEL if row['kind'] == SKIP else REVIEW_LABEL


def report_workbook(record):
    headers = ['Строка Excel', 'ID кавалера', 'ФИО', 'Результат', 'Причина', 'Биография', 'Источник', 'Проверка личности']
    rows = [[r['row'], r['person_id'], r['fio'], category(r), r['reason'], r['biography'], r['source_url'],
             '\n'.join(f'{name}: {value}' for name, value in r['evidence'].items() if value)] for r in record['rows']]
    return _workbook(headers, rows)


class ValidationError(ValueError):
    pass


def text(value):
    return '' if value is None else str(value).strip()


def blank(value):
    return value is None or isinstance(value, str) and not value.strip()


def key(value):
    return ' '.join(text(value).split()).casefold().replace('ё', 'е')


def _rank_key(value):
    value = key(value).removeprefix('гвардии ')
    for suffix in (' танковых войск', ' артиллерии', ' авиации'):
        value = value.removesuffix(suffix)
    return value


def canonical(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'))


def _state(settings):
    # Separate databases cannot share tokens or signatures. Never inside DB/media.
    identity = hashlib.sha256(str(settings.rewards_db_path.resolve()).encode()).hexdigest()
    root = Path(tempfile.gettempdir()) / 'fedorinov-biography-enrichment' / identity
    root.mkdir(parents=True, exist_ok=True, mode=0o700)
    return root


def _secret(settings):
    path = _state(settings) / 'signing.key'
    if path.exists():
        return path.read_bytes()
    fd, candidate = tempfile.mkstemp(dir=path.parent, suffix='.key')
    value = secrets.token_bytes(32)
    try:
        with os.fdopen(fd, 'wb') as stream:
            stream.write(value)
        try:
            os.link(candidate, path)  # Publish a complete key atomically.
        except FileExistsError:
            pass
    finally:
        Path(candidate).unlink(missing_ok=True)
    return path.read_bytes()


def _sign(settings, snapshot):
    return hmac.new(_secret(settings), snapshot.encode(), hashlib.sha256).hexdigest()


def _connection(path):
    connection = open_readonly_connection(path)
    connection.create_function('bio_blank', 1, lambda value: int(blank(value)), deterministic=True)
    connection.execute('begin')
    return connection


def _selection(selection):
    if not selection:
        return (), '', ()
    parts = selection.split('/')
    if len(parts) > 4 or any(not p.isascii() or not p.isdigit() for p in parts):
        raise ValidationError('Выберите награду или категорию из списка.')
    values = tuple(int(p) for p in parts)
    if any(p > 2**63 - 1 for p in values):
        raise ValidationError('Некорректная иерархия награды.')
    return values, ' and ' + ' and '.join(f'r.{name}=?' for name in IDS[:len(values)]), values


def statistics(settings):
    with closing(_connection(settings.rewards_db_path)) as connection:
        total, empty = connection.execute('select count(*), coalesce(sum(bio_blank(biography)),0) from person').fetchone()
        groups = []
        for depth in range(1, 5):
            ids = ','.join('r.' + name for name in IDS[:depth])
            joins = ' '.join(f'left join guide_lev_{i} g{i} on g{i}.id=r.{IDS[i]}' for i in range(depth))
            names = ','.join(f'g{i}.name' for i in range(depth))
            for row in connection.execute(f'''select {ids}, {names}, count(distinct p.id),
                    count(distinct case when bio_blank(p.biography)=1 then p.id end)
                    from rewards r join person p on p.id=r.person_id {joins} group by {ids}, {names}'''):
                # Null references stay visible as statistics but cannot be exported as a group.
                group_ids = row[:depth]
                if any(value is None for value in group_ids):
                    continue
                groups.append({'selection': '/'.join(map(str, group_ids)), 'depth': depth,
                               'label': ' / '.join(text(n) or 'Без названия' for n in row[depth:depth*2]),
                               'total': row[-2], 'empty': row[-1], 'filled': row[-2] - row[-1]})
        groups.sort(key=lambda item: (key(item['label']), item['selection']))
        return {'total': total, 'empty': empty, 'filled': total - empty, 'groups': groups}


def _identity(connection, person_id):
    row = connection.execute('''select p.id person_id,p.fio,p.birthday,p.id_rank,g.name rank
        from person p left join guide g on g.id=p.id_rank where p.id=?''', (person_id,)).fetchone()
    if row is None:
        return None
    result = dict(row)
    result['awards'] = [dict(r) for r in connection.execute('''select r.id,r.id_gos,r.id_catigory,
        r.id_sub_catigory,r.id_name,r.number,r.id_link,g0.name country,g1.name category,
        g2.name subcategory,g3.name award from rewards r
        left join guide_lev_0 g0 on g0.id=r.id_gos left join guide_lev_1 g1 on g1.id=r.id_catigory
        left join guide_lev_2 g2 on g2.id=r.id_sub_catigory left join guide_lev_3 g3 on g3.id=r.id_name
        where r.person_id=? order by r.id''', (person_id,))]
    return result


def _year(value):
    match = re.match(r'^(\d{4})(?:\D|$)', text(value))
    return match.group(1) if match else ''


def _original(snapshot):
    person = snapshot['person']
    selected = snapshot['selection']
    awards = [a for a in person['awards'] if all(a[IDS[i]] == v for i, v in enumerate(selected))]
    return {'person_id': str(person['person_id']), 'ФИО': text(person['fio']),
            'Год рождения': _year(person['birthday']), 'Звание': text(person['rank']),
            'Награда': ' | '.join(dict.fromkeys(text(a['award']) for a in awards)),
            'Иерархия ID': ' | '.join(dict.fromkeys(' / '.join(text(a[n]) for n in IDS) for a in awards)),
            'Номер(а) награды': ' | '.join(text(a['number']) for a in awards if a['number'] is not None)}


def _workbook(headers, rows, metadata=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Биографии'
    sheet.append(headers)
    for values in rows:
        if any(isinstance(value, str) and len(value) > 32000 for value in values):
            raise ValidationError('Слишком много сведений в одной записи. Выберите конкретную награду или меньшую группу.')
        sheet.append(values)
    for row in sheet:
        for cell in row:
            # User/database strings are literals, including leading '='.
            if isinstance(cell.value, str):
                cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='24364B')
    for index, name in enumerate(headers, 1):
        letter = sheet.cell(1, index).column_letter
        sheet.column_dimensions[letter].width = 64 if name in {'biography', 'snapshot', 'notes'} else 28
    sheet.freeze_panes = 'C2'
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 32
    if metadata is not None:
        for name, width in zip('ABCDEFG', (15, 36, 18, 45, 80, 55, 24)):
            sheet.column_dimensions[name].width = width
        sheet.freeze_panes = 'C2'
        status = DataValidation(type='list', formula1='"Найдено,Требует проверки"', allow_blank=True)
        status.errorTitle = 'Выберите статус'
        status.error = 'Допустимы «Найдено» и «Требует проверки».'
        status.showErrorMessage = True
        status.errorStyle = 'stop'
        status.promptTitle = 'Подтверждение личности'
        status.prompt = '«Найдено»: источник подтверждает ФИО, достоверный год рождения и награду.'
        status.showInputMessage = True
        sheet.add_data_validation(status)
        status.add(f'G2:G{len(rows)+1}')
        hidden = workbook.create_sheet(METADATA_SHEET)
        hidden.append(HEADERS)
        for values in metadata:
            if any(isinstance(value, str) and len(value) > 32000 for value in values):
                raise ValidationError('Слишком много наград в одной записи. Выберите более узкую группу.')
            hidden.append(values)
        for row in hidden:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        hidden.sheet_state = 'hidden'
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    content = stream.getvalue()
    with ZipFile(BytesIO(content)) as archive:
        if len(content) > MAX_BYTES or sum(item.file_size for item in archive.infolist()) > MAX_UNPACKED:
            raise ValidationError('Выгрузка слишком велика. Уменьшите количество записей.')
    return content


def export_workbook(settings, selection='', limit=20):
    if not 1 <= limit <= MAX_ROWS:
        raise ValidationError(f'Количество записей: от 1 до {MAX_ROWS}.')
    selected, clause, params = _selection(selection)
    rows = []
    metadata = []
    with closing(_connection(settings.rewards_db_path)) as connection:
        eligible = connection.execute('''select p.id from person p where bio_blank(p.biography)=1'''
            + (f' and exists(select 1 from rewards r where r.person_id=p.id {clause})' if selected else '')
            + ' order by p.id limit ?', (*params, limit)).fetchall()
        for row in eligible:
            snapshot = {'version': 2, 'contract': 'visible-7-v2', 'selection': selected,
                        'person': _identity(connection, row[0])}
            original = _original(snapshot)
            serialized = canonical(snapshot)
            original.update({'snapshot': serialized, 'snapshot_signature': _sign(settings, serialized),
                             'Примечание к году': '1945: возможное значение по умолчанию; требуется ручная проверка'
                             if original['Год рождения'] == '1945' else ''})
            metadata.append([original.get(header, '') for header in HEADERS])
            rows.append([original['person_id'], original['ФИО'], original['Год рождения'], original['Награда'], '', '', ''])
    if not rows:
        raise ValidationError('В выбранной группе нет пустых биографий для экспорта.')
    return _workbook(VISIBLE_HEADERS, rows, metadata)


def _sheet_rows(sheet, required, *, human=False):
    sheet.reset_dimensions()
    iterator = sheet.iter_rows()
    headers = None
    for header_row in range(1, 6):
        cells = next(iterator, ())
        names = [ALIASES.get(text(c.value), text(c.value)) for c in cells]
        if 'person_id' in names and 'ФИО' in names:
            if any(c.data_type in {'f', 'e'} for c in cells):
                raise ValidationError('Заголовки Excel содержат формулы или ошибки.')
            if len(names) > MAX_COLUMNS or len([n for n in names if n]) != len(set(n for n in names if n)):
                raise ValidationError('Повторяющиеся заголовки или слишком много колонок.')
            headers = names
            break
    if headers is None or not required.issubset(headers):
        raise ValidationError('Не найдены обязательные колонки. Скачайте Excel в разделе «Биографии» и сохраните его структуру.')
    if human and set(headers) != required:
        raise ValidationError('На основном листе должно быть семь колонок исходной выгрузки. Скачайте новый Excel.')
    rows = []
    for number, cells in enumerate(iterator, header_row + 1):
        if number - header_row > MAX_ROWS or len(cells) > MAX_COLUMNS:
            raise ValidationError('Допустимо не более 5 000 строк и 40 колонок.')
        if human and len(cells) > len(headers):
            raise ValidationError('В строке Excel есть лишние колонки без заголовков. Сохраните семь колонок исходной выгрузки.')
        if not any(text(c.value) for c in cells):
            continue
        values = {name: text(cells[i].value) if i < len(cells) else '' for i, name in enumerate(headers) if name}
        if 'biography' in headers:
            bio_index = headers.index('biography')
            if bio_index < len(cells) and isinstance(cells[bio_index].value, str):
                values['biography'] = cells[bio_index].value
        if any(len(text(c.value)) > 32000 for c in cells):
            raise ValidationError('Слишком длинная ячейка Excel.')
        rows.append({'row': number, 'values': values, 'invalid': any(c.data_type in {'f', 'e'} for c in cells)})
    if not rows:
        raise ValidationError('В файле нет записей.')
    return rows


def read_rows(content):
    if not content or len(content) > MAX_BYTES:
        raise ValidationError('Выберите XLSX размером не более 10 МБ.')
    workbook = None
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 200 or sum(item.file_size for item in entries) > MAX_UNPACKED:
                raise ValidationError('Распакованный XLSX превышает допустимый размер 50 МБ.')
            if any('vbaProject' in item.filename for item in entries):
                raise ValidationError('XLSX с макросами не поддерживается.')
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        sheet = workbook.worksheets[0]
        first = next(sheet.iter_rows(min_row=1, max_row=1), ())
        human = 'ID кавалера' in [text(c.value) for c in first]
        if human:
            required = {ALIASES.get(name, name) for name in VISIBLE_HEADERS}
            rows = _sheet_rows(sheet, required, human=True)
            if METADATA_SHEET not in workbook.sheetnames:
                raise ValidationError('В файле нет служебного листа исходной выгрузки. Скачайте Excel заново; не переносите только видимый лист.')
            meta_rows = _sheet_rows(workbook[METADATA_SHEET], {'person_id', 'ФИО', 'Звание', 'Иерархия ID',
                                                              'Номер(а) награды', 'snapshot', 'snapshot_signature'})
            metadata = {}
            for entry in meta_rows:
                values = entry['values']
                if entry['invalid'] or values['person_id'] in metadata:
                    raise ValidationError('Служебные данные повреждены: формулы или повторяющиеся ID. Скачайте новую выгрузку.')
                metadata[values['person_id']] = values
            for entry in rows:
                visible = entry['values']
                values = dict(metadata.get(visible['person_id'], {}))
                values.update(visible)
                values['_format'] = 'visible-7-v2'
                entry['values'] = values
            return rows
        required = {'person_id', 'ФИО', 'Год рождения', 'Звание', 'Награда', 'Иерархия ID',
                    'Номер(а) награды', 'biography', 'source_url', 'status'}
        return _sheet_rows(sheet, required)
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError('Не удалось прочитать XLSX. Файл повреждён или имеет неподдерживаемый формат.') from exc
    finally:
        if workbook is not None:
            workbook.close()


def _snapshot(settings, values):
    serialized = values.get('snapshot', '')
    signature = values.get('snapshot_signature', '')
    if not serialized or not hmac.compare_digest(signature, _sign(settings, serialized)):
        raise ValidationError('Исходные служебные данные отсутствуют или изменены. Приложите исходный Excel с теми же кавалерами либо скачайте новую выгрузку и перенесите только найденные биографии, источники и статусы.')
    try:
        snapshot = json.loads(serialized)
        if snapshot['version'] not in {1, 2} or not isinstance(snapshot['person']['person_id'], int):
            raise ValueError()
        return snapshot
    except (ValueError, KeyError, TypeError) as exc:
        raise ValidationError('Повреждены служебные данные исходной выгрузки. Скачайте Excel заново.') from exc


def _assess(connection, settings, entry, originals, duplicates):
    values = entry['values']
    human = values.get('_format') == 'visible-7-v2'
    result = {'row': entry['row'], 'person_id': values.get('person_id', ''), 'fio': values.get('ФИО', ''),
              'biography': values.get('biography', ''), 'source_url': values.get('source_url', ''),
              'evidence': {label: values.get(name, '') for name, label in (
                  ('ФИО в карточке', 'ФИО в источнике'), ('Год в карточке', 'Год в источнике'),
                  ('Награда в карточке', 'Награда в источнике'), ('Звание в карточке', 'Звание в источнике'),
                  ('Идентификация', 'Проверка личности'), ('notes', 'Подтверждённые факты'), ('status', 'Статус в файле'))},
              'kind': ERROR, 'reason': ''}
    if human:
        result['evidence'] = {'ФИО в файле': values.get('ФИО', ''), 'Год рождения в файле': values.get('Год рождения', ''),
                              'Награда в файле': values.get('Награда', ''), 'Статус в файле': values.get('status', '')}
    def fail(kind, reason):
        result.update(kind=kind, reason=reason)
        return result
    if entry['invalid']:
        return fail(ERROR, 'Формулы и ошибки Excel: замените значениями.')
    raw_id = values['person_id']
    if not re.fullmatch(r'[1-9][0-9]{0,18}', raw_id) or int(raw_id) > 2**63 - 1:
        return fail(ERROR, 'Некорректный ID кавалера. Сохраните ID из исходного Excel без изменений.')
    person_id = int(raw_id)
    if person_id in duplicates:
        return fail(CONFLICT, 'ID кавалера повторяется: все его строки требуют проверки. Оставьте одну согласованную запись.')
    current = _identity(connection, person_id)
    if current is None:
        return fail(SKIP, 'Кавалер с таким ID отсутствует в текущей базе. Новые записи не создаются.')
    try:
        if human and not values.get('snapshot'):
            raise ValidationError('Кавалер не входит в исходную выгрузку этого файла. Скачайте Excel для нужной группы; не заменяйте ID вручную.')
        source = values if values.get('snapshot') else originals.get(str(person_id), {})
        snapshot = _snapshot(settings, source)
        if human and (snapshot['version'] != 2 or snapshot.get('contract') != 'visible-7-v2'):
            raise ValidationError('Служебные данные относятся к другому формату выгрузки. Скачайте новый Excel в этом разделе.')
    except ValidationError as exc:
        return fail(CONFLICT, str(exc))
    if snapshot['person'] != current:
        return fail(CONFLICT, 'Исходные данные кавалера или его награды изменились после выгрузки.')
    original = _original(snapshot)
    for name, expected in original.items():
        if key(values.get(name)) != key(expected):
            label = 'ID кавалера' if name == 'person_id' else name
            given = text(values.get(name)) or 'не заполнено'
            return fail(CONFLICT, f'Не совпадает «{label}»: в файле «{given}», в исходной выгрузке «{expected}». Проверьте запись и источник.')
    result['snapshot'] = snapshot
    biography = connection.execute('select biography from person where id=?', (person_id,)).fetchone()[0]
    if not blank(biography):
        return fail(SKIP, 'Биография уже заполнена; существующий текст сохранён.')
    status = key(values['status'])
    confirmed = {'найдено'} if human else {'биография подтверждена', 'подтверждено', 'confirmed'}
    if 'требует проверки' in status or status in {'requires review', 'requires verification'}:
        return fail(CONFLICT, 'В файле указан статус «Требует проверки». Уточните личность и источник; строка не будет сохранена.')
    if status and status not in confirmed:
        return fail(CONFLICT, f'Неизвестный статус «{values["status"]}». Используйте «Найдено» или «Требует проверки» в новой выгрузке.')
    if not original['Год рождения'] or original['Год рождения'] == '1945':
        return fail(CONFLICT, 'Год рождения не подтверждён: он отсутствует либо 1945 мог быть подставлен при старом импорте. Требуется ручная проверка.')
    if blank(result['biography']):
        return fail(SKIP, 'Биография в файле не заполнена. Записывать нечего.')
    if status not in confirmed:
        return fail(CONFLICT, 'Статус не заполнен. Укажите «Найдено» только после подтверждения личности по ФИО, году рождения и награде.')
    if not human and key(values.get('ФИО в карточке')) != key(current['fio']):
        return fail(CONFLICT, 'ФИО в карточке источника не совпадает с исходным.')
    if not human and text(values.get('Год в карточке')) != original['Год рождения']:
        return fail(CONFLICT, 'Год рождения в карточке источника не совпадает.')
    # ALE-491 has no separate source-award column: its accepted three-signal
    # identity statement is an explicit mapping, never inferred from biography.
    source_award = values.get('Награда в карточке', '')
    if not human and source_award:
        if key(source_award) not in {key(a) for a in original['Награда'].split(' | ')}:
            return fail(CONFLICT, 'Награда в карточке источника не совпадает.')
    elif not human and (values.get('snapshot') or key(values.get('Идентификация')) != 'фио, реальный год и награда совпали'):
        return fail(CONFLICT, 'В прежнем файле нет подтверждения совпадения ФИО, реального года и награды. Уточните сведения или используйте новую выгрузку.')
    if not human and (not text(values.get('notes')) or not text(values.get('Идентификация'))):
        return fail(CONFLICT, 'В прежнем файле не заполнены проверка личности и подтверждённые факты. Уточните сведения или используйте новую выгрузку.')
    if not human and not blank(values.get('Звание в карточке')) and _rank_key(values['Звание в карточке']) != _rank_key(current['rank']):
        # Pilot profiles can give a later service rank. Require its explicit
        # source comparison, rather than silently treating any rank as equal.
        comparison = key(values.get('Сверка звания'))
        pilot_comparison = (not values.get('snapshot') and
            key(values.get('Идентификация')) == 'фио, реальный год и награда совпали' and
            comparison.startswith('звание sqlite отличается от вводной части карточки;') and
            f'sqlite: {key(current["rank"])};' in comparison and
            f'warheroes (вводная часть): {key(values["Звание в карточке"])}.' in comparison)
        if not pilot_comparison:
            return fail(CONFLICT, 'Звание источника отличается; требуется ручная проверка.')
    urls = values['source_url'].splitlines()
    try:
        valid_urls = bool(urls) and all(urlsplit(url).scheme in {'https', 'http'} and urlsplit(url).hostname
            and not urlsplit(url).username and not urlsplit(url).password for url in urls)
    except ValueError:
        valid_urls = False
    if not valid_urls:
        return fail(ERROR, 'Укажите полный http/https URL карточки источника.')
    if len(result['biography']) > 2500:
        return fail(ERROR, 'Биография длиннее 2 500 Unicode символов; текст не обрезан.')
    return fail(READY, 'Исходные ID, ФИО, год рождения и награды совпали; биография в базе пуста. Личность подтверждена в файле.' +
                (' Различие званий явно пояснено в прежнем файле; звание в базе не меняется.'
                 if key(values.get('Звание в карточке')) and key(values['Звание в карточке']) != key(current['rank']) else ''))


def _save(settings, token, record):
    root = _state(settings)
    fd, path = tempfile.mkstemp(dir=root, suffix='.tmp')
    try:
        with os.fdopen(fd, 'w', encoding='utf-8') as stream:
            stream.write(canonical(record))
        os.replace(path, root / (token + '.json'))
    finally:
        Path(path).unlink(missing_ok=True)


def load_preview(settings, token):
    if not re.fullmatch(r'[0-9a-f]{64}', token):
        raise ValidationError('Предпросмотр не найден. Загрузите XLSX заново.')
    try:
        record = json.loads((_state(settings) / (token + '.json')).read_text(encoding='utf-8'))
    except (OSError, ValueError) as exc:
        raise ValidationError('Предпросмотр не найден. Загрузите XLSX заново.') from exc
    if time.time() - record['created'] > TTL:
        raise ValidationError('Прошло более 7 дней. Загрузите файл для нового предпросмотра.')
    return record


def preview(settings, content, original_content=None):
    entries = read_rows(content)
    originals = {}
    if original_content:
        for entry in read_rows(original_content):
            values = entry['values']
            if entry['invalid'] or values['person_id'] in originals:
                raise ValidationError('Исходная выгрузка содержит формулы или повторяющиеся ID.')
            _snapshot(settings, values)
            originals[values['person_id']] = values
    counts = Counter(int(e['values']['person_id']) for e in entries if re.fullmatch(r'[1-9][0-9]{0,18}', e['values']['person_id']))
    duplicates = {value for value, count in counts.items() if count > 1}
    with closing(_connection(settings.rewards_db_path)) as connection:
        rows = [_assess(connection, settings, entry, originals, duplicates) for entry in entries]
    token = secrets.token_hex(32)
    record = {'created': time.time(), 'state': 'preview', 'file_sha256': hashlib.sha256(content).hexdigest(),
              'entries': entries, 'originals': originals, 'duplicates': list(duplicates), 'rows': rows,
              'ready': sum(r['kind'] == READY for r in rows), 'saved': 0}
    _save(settings, token, record)
    return token, record


def commit(settings, token):
    ensure_write_allowed(settings)
    with closing(open_write_connection(settings.rewards_db_path, settings.write_mode)) as connection:
        connection.execute('begin immediate')
        try:
            record = load_preview(settings, token)
            if record['state'] == 'committed':
                return record
            if record['state'] != 'preview':
                raise ValidationError('Результат предыдущей попытки не подтверждён. Загрузите Excel для новой проверки.')
            # Only rows explicitly shown as ready may be applied. New eligibility
            # after preview never expands the owner's confirmed subset.
            eligible_rows = {r['row'] for r in record['rows'] if r['kind'] == READY}
            checked = [_assess(connection, settings, entry, record['originals'], set(record['duplicates']))
                       for entry in record['entries']]
            saved = 0
            for row in checked:
                if row['row'] not in eligible_rows or row['kind'] != READY:
                    if row['kind'] == READY:
                        row.update(kind=SKIP, reason='Строка не входила в подтверждённую группу готовых записей. Загрузите файл для новой проверки.')
                    continue
                # The reserved SQLite lock covers revalidation and every UPDATE.
                current = connection.execute('select biography from person where id=?', (int(row['person_id']),)).fetchone()
                if current is None or not blank(current[0]):
                    raise ValidationError('Запись изменена другим пользователем: сохранение отменено. Загрузите файл для новой проверки.')
                cursor = connection.execute('update person set biography=? where id=?', (row['biography'], int(row['person_id'])))
                if cursor.rowcount != 1:
                    raise ValidationError('База данных отклонила запись. Все изменения отменены; загрузите файл для новой проверки.')
                row.update(kind='Сохранено', reason='Сохранена только ранее пустая биография.')
                saved += 1
            record.update(rows=checked, saved=saved, ready=0, state='commit_pending')
            # Preserve evidence before committing SQLite. On a process crash this
            # state is deliberately uncertain; replay requires a fresh preview.
            _save(settings, token, record)
            connection.commit()
        except BaseException:
            connection.rollback()
            raise
    record['state'] = 'committed'
    try:
        _save(settings, token, record)
    except OSError:
        # Data is committed. Never claim rollback for a report storage failure.
        record['report_warning'] = 'Данные сохранены, но окончательный статус отчёта не удалось записать. Повторная проверка файла безопасна.'
    return record
