"""XLSX import: one selected reference, number-only deduplication, no media writes."""
from contextlib import closing
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
import sqlite3
from zipfile import ZipFile

from openpyxl import load_workbook

from ..config import Settings
from ..db import open_write_connection
from ..repositories.persons_write import person_data_from_mapping
from ..repositories.reward_reference import reward_reference_from_connection
from ..repositories.rewards_write import reward_data_from_mapping
from .audit import log_action
from .dates import normalize_birth_year_input, normalize_date_input
from .write_guard import ensure_write_allowed

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_UNPACKED_BYTES = 100 * 1024 * 1024
MAX_ROWS = 50000
MAX_COLUMNS = 64
HEADERS = {
    'Фамилия': ('фамилия',),
    'Имя': ('имя',),
    'Отчество': ('отчество',),
    'Номер ордена': ('номер ордена', 'номер награды', 'номер выбранной награды', 'номер'),
    'Год рождения': ('год рождения', 'дата рождения', 'дата/год рождения', 'дата / год рождения'),
    'Звание': ('звание', 'звание / специальность', 'звание/специальность', 'специальность'),
}
REQUIRED_HEADERS = ('Фамилия', 'Имя', 'Отчество', 'Номер ордена')


class ImportValidationError(ValueError):
    pass


@dataclass
class ImportResult:
    added: int = 0
    skipped: int = 0
    errors: list[dict[str, object]] = field(default_factory=list)

    def fail(self, row: int, reason: str) -> None:
        self.errors.append({'row': row, 'reason': reason})


def _text(value: object) -> str:
    return '' if value is None else str(value).strip()


def _key(value: object) -> str:
    return ' '.join(_text(value).split()).casefold()


def _read_rows(content: bytes) -> list[tuple[int, dict[str, object]]]:
    if not content or len(content) > MAX_UPLOAD_BYTES:
        raise ImportValidationError('Выберите XLSX размером не более 10 МБ.')
    workbook = None
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_UNPACKED_BYTES:
                raise ImportValidationError('Распакованный XLSX слишком большой (не более 100 МБ).')
        # Keep formulas visible: cached results must not silently become imported data.
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        if not workbook.worksheets:
            raise ImportValidationError('В XLSX нет листа с данными.')
        sheet = workbook.worksheets[0]
        # Do not trust producer-provided dimensions (some exporters incorrectly write A1:A1).
        sheet.reset_dimensions()
        rows = sheet.iter_rows()
        header = next(rows, ())
        if len(header) > MAX_COLUMNS:
            raise ImportValidationError('В XLSX слишком много колонок (не более 64).')
        columns = {}
        for index, cell in enumerate(header):
            for name, aliases in HEADERS.items():
                if _key(cell.value) in aliases:
                    if name in columns:
                        raise ImportValidationError(f'Колонка «{name}» указана несколько раз.')
                    columns[name] = index
        missing = [name for name in REQUIRED_HEADERS if name not in columns]
        if missing:
            raise ImportValidationError('Не найдены обязательные колонки: ' + ', '.join(missing) + '.')
        parsed = []
        for row_number, cells in enumerate(rows, start=2):
            if row_number > MAX_ROWS + 1 or len(cells) > MAX_COLUMNS:
                raise ImportValidationError('Допустимо не более 50 000 строк и 64 колонок.')
            if not any(_text(cell.value) for cell in cells):
                continue
            values = {name: cells[index].value if index < len(cells) else None for name, index in columns.items()}
            values['_invalid_cell'] = any(
                cells[index].data_type in {'f', 'e'} for index in columns.values() if index < len(cells)
            )
            parsed.append((row_number, values))
        if not parsed:
            raise ImportValidationError('В первом листе XLSX нет строк для импорта.')
        return parsed
    except ImportValidationError:
        raise
    except Exception as exc:
        raise ImportValidationError('Не удалось прочитать XLSX. Проверьте файл и сохраните его заново в Excel.') from exc
    finally:
        if workbook is not None:
            workbook.close()


def _birth_year(value: object) -> str:
    if not _text(value):
        return '1945'
    if isinstance(value, (datetime, date)):
        year = str(value.year)
    else:
        year = _text(value)
        if isinstance(value, float) and value.is_integer():
            year = str(int(value))
        if len(year) != 4 or not year.isascii() or not year.isdigit():
            year = normalize_date_input(year, required=True)[:4]
    return normalize_birth_year_input(year, required=True)


def _number(value: object) -> int:
    if isinstance(value, bool) or not _text(value):
        raise ImportValidationError('Укажите номер награды.')
    if isinstance(value, float):
        if not value.is_integer():
            raise ImportValidationError('Номер награды должен быть целым числом.')
        value = int(value)
    number = reward_data_from_mapping({'number': value}).number
    if not -(2**63) <= number < 2**63:
        raise ImportValidationError('Номер награды выходит за допустимый диапазон.')
    return number


def import_persons(settings: Settings, name_id: int, content: bytes) -> ImportResult:
    ensure_write_allowed(settings)
    rows = _read_rows(content)  # Exhaust/validate workbook before opening a write transaction.
    result = ImportResult()
    with closing(open_write_connection(settings.rewards_db_path, settings.write_mode)) as connection:
        # Serialize with manual writers and concurrent imports before taking the number snapshot.
        connection.execute('begin immediate')
        try:
            reference = reward_reference_from_connection(connection, name_id)
            if reference is None:
                raise ImportValidationError('Выбранная награда отсутствует в справочнике.')
            ranks = {}
            for row in connection.execute('select id, name from guide'):
                ranks.setdefault(_key(row['name']), []).append(row['id'])
            fallback = ranks.get(_key('Красноармеец'), [])
            if len(fallback) != 1:
                raise ImportValidationError('Импорт остановлен: в справочнике нет единственного значения «Красноармеец». Проверьте справочник до массовой записи.')
            # Exactly the same trim(cast(number as text)) semantics as _find_reward_duplicate.
            # One scan per batch, never one full-table scan per incoming row.
            numbers = {row[0] for row in connection.execute(
                'select trim(cast(number as text)) from rewards where id_name = ?', (name_id,)
            )}
            prepared = []
            for row_number, values in rows:
                number = None
                try:
                    number = _number(values['Номер ордена'])
                    if str(number) in numbers:
                        result.skipped += 1
                        continue
                    if values['_invalid_cell']:
                        raise ImportValidationError('Формулы и ошибки Excel не поддерживаются: замените их значениями.')
                    for name in ('Фамилия', 'Имя'):
                        if not isinstance(values[name], str) or not values[name].strip():
                            raise ImportValidationError(f'Заполните колонку «{name}» текстом.')
                    if values['Отчество'] is not None and not isinstance(values['Отчество'], str):
                        raise ImportValidationError('Колонка «Отчество» должна содержать текст.')
                    rank = _key(values.get('Звание'))
                    rank_ids = ranks.get(rank, []) if rank else fallback
                    if len(rank_ids) != 1:
                        raise ImportValidationError('Звание не найдено однозначно в существующем справочнике: ' + _text(values.get('Звание')))
                    person = person_data_from_mapping({
                        'fio': ' '.join(_text(values[name]) for name in ('Фамилия', 'Имя', 'Отчество') if _text(values[name])),
                        'birthday': _birth_year(values.get('Год рождения')),
                        'id_rank': rank_ids[0],
                    })
                    prepared.append((row_number, person, number, None))
                except ValueError as exc:
                    prepared.append((row_number, None, number, str(exc)))
            # All preflight checks above finish before the first INSERT. Each person/reward pair is atomic.
            for row_number, person, number, error in prepared:
                if number is not None and str(number) in numbers:
                    result.skipped += 1
                    continue
                if error is not None:
                    result.fail(row_number, error)
                    continue
                connection.execute('savepoint import_row')
                try:
                    cursor = connection.execute(
                        'insert into person (fio, birthday, id_rank) values (?, ?, ?)',
                        (person.fio, person.birthday, person.id_rank),
                    )
                    connection.execute(
                        '''insert into rewards (person_id, id_gos, id_catigory, id_sub_catigory, id_name, id_link, number, instock)
                           values (?, ?, ?, ?, ?, ?, ?, ?)''',
                        (cursor.lastrowid, reference['id_gos'], reference['id_catigory'], reference['id_sub_catigory'],
                         name_id, reference['id_link'], number, False),
                    )
                except sqlite3.IntegrityError:
                    connection.execute('rollback to import_row')
                    result.fail(row_number, 'Запись отклонена ограничениями базы; кавалер и награда не сохранены.')
                else:
                    numbers.add(str(number))
                    result.added += 1
                finally:
                    connection.execute('release import_row')
            connection.commit()
        except Exception:
            connection.rollback()
            raise
    log_action('import', 'person', details={'id_name': name_id, 'added': result.added, 'skipped': result.skipped, 'errors': len(result.errors)})
    return result
