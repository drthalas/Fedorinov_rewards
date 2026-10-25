"""Owner FAIL regression coverage; every database/media write is TEMP only."""
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing
from io import BytesIO
from unittest.mock import patch
from zipfile import ZipFile,ZIP_DEFLATED
import copy,json,sqlite3,unittest

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app.main import app
from backend.app.services import biography_enrichment as bio, manual_import_review as manual, manual_review_choices as choices, enrichment_workflow as flow
from backend.app.services.write_guard import WriteBlockedError
from tests import test_ale513_manual_import_review as fixtures


class ImportWorkflowTests(unittest.TestCase):
    setUp=fixtures.ManualReviewTests.setUp
    execute=fixtures.ManualReviewTests.execute
    query=fixtures.ManualReviewTests.query

    def content(self, *, year='1930', biography='Проверенная биография.', fio=None,
                identity='ФИО, точная I степень и номер совпали', award='Орден Трудовой Славы I степени',
                source='https://warheroes.ru/hero/hero.asp?Hero_id=1', status='Требует проверки'):
        entries=bio.read_rows(bio.export_workbook(self.settings,'',20));rows=[]
        for entry in entries:
            v=entry['values'];v.update({'biography':biography,'source_url':source,'status':status,
                'Год в карточке':year,'ФИО в карточке':fio or v['ФИО'],'Награда в карточке':award,
                'Идентификация':identity,'notes':f'В источнике указаны полное ФИО, {award} №18 и год рождения {year}.'})
            rows.append([v.get(h,'') for h in bio.HEADERS])
        return bio._workbook(bio.HEADERS,rows)

    def preview(self, *, current='1945', **kwargs):
        self.execute(f"update person set birthday='{current}';update rewards set number=18;")
        self.last_content=self.content(**kwargs)
        return bio.preview(self.settings,self.last_content)

    def import_plan(self, token):
        return flow.prepare_import(self.settings,token,flow.present(self.settings,token)['revision'])

    def choose(self, token, **values):
        return choices.prepare(self.settings,token,2,{'review_revision':choices.detail(self.settings,token,2)['revision'],**values})

    def test_verified_1945_and_empty_biography_are_automatic_only_on_explicit_import(self):
        token,record=self.preview()
        snapshot=record['entries'][0]['values']['snapshot'];before=self.db.read_bytes()
        detail=choices.detail(self.settings,token,2)
        self.assertEqual(detail['fields'],{})
        self.assertEqual(set(detail['automatic']),{'birthday','biography'})
        op,plan=self.import_plan(token)
        self.assertEqual(self.db.read_bytes(),before)
        self.assertEqual(set(plan['rows'][0]['writes']),{'birthday','biography'})
        result=flow.commit_import(self.settings,op)
        self.assertEqual((result['saved'],result['fields_saved']),(1,2))
        self.assertEqual(self.query('select birthday,biography from person'),[('1930','Проверенная биография.')])
        self.assertEqual(bio.load_preview(self.settings,token)['entries'][0]['values']['snapshot'],snapshot)
        view=flow.present(self.settings,token)
        self.assertEqual(view['rows'][0]['state'],'confirmed')
        self.assertEqual(view['ready'],0)
        self.assertEqual(flow.commit_import(self.settings,op),result)
        self.assertEqual(self.query('select count(*) from logs where action like "manual_import_review:%"'),[(1,)])
        with self.assertRaises(bio.ValidationError):self.import_plan(token)

    def test_biography_only_never_has_blank_vs_filled_chooser(self):
        token,_=self.preview(current='1918',year='1918')
        detail=choices.detail(self.settings,token,2)
        self.assertEqual(detail['fields'],{})
        self.assertEqual(set(detail['automatic']),{'biography'})
        with TestClient(app) as client:
            response=client.get(f'/data-maintenance/enrichment/biographies/review/{token}/2')
            self.assertNotIn('type="radio"',response.text)
            self.assertIn('Прочитать биографию',response.text)
            self.assertIn('Будет импортировано',response.text)

    def test_partial_biography_import_then_only_remaining_non1945_year(self):
        token,_=self.preview(current='1932',year='1931')
        op,_=self.import_plan(token);flow.commit_import(self.settings,op)
        detail=choices.detail(self.settings,token,2)
        self.assertEqual(set(detail['fields']),{'birthday'})
        self.assertEqual(detail['automatic'],{})
        self.assertEqual(detail['status'],'Частично подтверждено')
        p,_=self.choose(token,birthday_choice='proposed')
        event=flow.commit_review(self.settings,p)
        self.assertEqual(event['writes'],['birthday'])
        self.assertEqual(self.query('select birthday from person'),[('1931',)])
        self.assertTrue(choices.detail(self.settings,token,2)['confirmed'])
        # Re-uploading identical bytes also uses authenticated prior results.
        again,_=bio.preview(self.settings,self.last_content)
        self.assertEqual(flow.present(self.settings,again)['rows'][0]['state'],'confirmed')

    def test_rejecting_non1945_year_resolves_it_without_repeat_review(self):
        token,_=self.preview(current='1932',year='1931',biography='')
        p,_=self.choose(token,birthday_choice='current')
        q,_=self.choose(token,birthday_choice='proposed')
        flow.commit_review(self.settings,p)
        self.assertTrue(choices.detail(self.settings,token,2)['confirmed'])
        with self.assertRaises(bio.ValidationError):flow.commit_review(self.settings,q)
        self.assertEqual(self.query('select birthday from person'),[('1932',)])
        history=flow.history_view(manual.durable_history(self.settings,11095))
        self.assertEqual(history[0]['fields'][0]['decision'],'Отклонено предложение')
        self.assertEqual(history[0]['fields'][0]['proposed'],'1931')
        self.assertEqual(history[0]['fields'][0]['after'],'1932')

    def test_genuine_1945_is_unchanged_and_requires_no_pseudo_choice(self):
        token,_=self.preview(current='1945-05-09',year='1945')
        detail=choices.detail(self.settings,token,2)
        self.assertEqual(detail['fields'],{})
        self.assertEqual(set(detail['automatic']),{'biography'})
        op,_=self.import_plan(token);flow.commit_import(self.settings,op)
        self.assertEqual(self.query('select birthday from person'),[('1945-05-09',)])
        self.assertTrue(choices.detail(self.settings,token,2)['confirmed'])

    def test_ermakova_unresolved_identity_has_explanation_and_no_inactive_options(self):
        token,_=self.preview(current='1954-05-09',year='1954',fio='Ермакова Мария Владимировна',identity='Личность не подтверждена')
        detail=choices.detail(self.settings,token,2)
        self.assertEqual(detail['state'],'blocked')
        self.assertFalse(detail['automatic']);self.assertFalse(detail['fields'])
        with TestClient(app) as client:
            response=client.get(f'/data-maintenance/enrichment/biographies/review/{token}/2')
            self.assertIn('В базе:',response.text);self.assertIn('В источнике:',response.text)
            self.assertNotIn('type="radio"',response.text)
            self.assertNotIn('data-review-save',response.text)
        with self.assertRaises(bio.ValidationError):self.import_plan(token)

    def test_full_fio_award_degree_serial_and_source_required_for_auto(self):
        cases=[{'fio':'Другой Иван Иванович'},{'award':'Трудовая Слава II степени'},
               {'source':'https://example.org/card'},{'identity':'Личность не подтверждена'},{'identity':'Требует проверки'}]
        for args in cases:
            token,_=self.preview(**args)
            self.assertNotIn('birthday',choices.detail(self.settings,token,2)['automatic'])
        token,_=self.preview();record=bio.load_preview(self.settings,token)
        record['entries'][0]['values']['notes']='Источник указывает награду №19; год1930.'
        bio._save(self.settings,token,record)
        self.assertFalse(choices.detail(self.settings,token,2)['automatic'])

    def test_existing_biography_is_never_replaced(self):
        self.execute("update person set birthday='1945';update rewards set number=18")
        content=self.content();self.execute("update person set biography='Существующий текст'")
        token,_=bio.preview(self.settings,content)
        detail=choices.detail(self.settings,token,2)
        self.assertEqual(set(detail['automatic']),{'birthday'})
        self.assertNotIn('biography',detail['fields'])
        op,_=self.import_plan(token);flow.commit_import(self.settings,op)
        self.assertEqual(self.query('select biography from person'),[('Существующий текст',)])

    def test_stale_concurrent_and_bad_receipts_fail_closed(self):
        token,_=self.preview();op,_=self.import_plan(token)
        self.execute("update person set biography='Параллельная правка'")
        with self.assertRaises(bio.ValidationError):flow.commit_import(self.settings,op)
        self.assertEqual(self.query('select birthday from person'),[('1945',)])
        self.execute('update person set biography=null')
        token,_=self.preview();op,_=self.import_plan(token);flow.commit_import(self.settings,op)
        self.execute("update person set birthday='1935'")
        self.assertEqual(flow.present(self.settings,token)['rows'][0]['state'],'blocked')
        self.execute("update person set birthday='1930'")
        event=json.loads(self.query('select comment from logs where action like "manual_import_review:%"')[0][0])
        event['fields']['birthday']['after']='1935'
        with closing(sqlite3.connect(self.db)) as c:c.execute('update logs set comment=? where action like "manual_import_review:%"',(json.dumps(event),));c.commit()
        self.assertEqual(flow.present(self.settings,token)['rows'][0]['state'],'blocked')

    def test_batch_rollback_and_concurrent_retry_are_atomic(self):
        token,_=self.preview();op,_=self.import_plan(token)
        self.execute("create trigger fail_audit before insert on logs begin select raise(abort,'TEMP audit failure');end;")
        with self.assertRaises(sqlite3.IntegrityError):flow.commit_import(self.settings,op)
        self.assertEqual(self.query('select birthday,biography from person'),[('1945',None)])
        self.assertEqual(self.query('select count(*) from logs'),[(0,)])
        self.execute('drop trigger fail_audit')
        with ThreadPoolExecutor(2) as pool:results=list(pool.map(lambda _:flow.commit_import(self.settings,op),range(2)))
        self.assertEqual(results[0],results[1])
        self.assertEqual(self.query('select count(*) from logs where action like "manual_import_review:%"'),[(1,)])

    def test_bad_hmac_wrong_database_and_unsigned_input_cannot_be_auto_imported(self):
        token,_=self.preview();record=bio.load_preview(self.settings,token)
        record['entries'][0]['values']['snapshot_signature']='0'*64;bio._save(self.settings,token,record)
        with self.assertRaises(bio.ValidationError):self.import_plan(token)
        content=self.content()
        other=self.root/'other';other.mkdir()
        with patch.object(bio,'_state',return_value=other):
            token,_=bio.preview(self.settings,content)
            with self.assertRaises(bio.ValidationError):self.import_plan(token)
        token,_=self.preview();record=bio.load_preview(self.settings,token)
        record['entries'][0]['values']['snapshot']='';record['entries'][0]['values']['snapshot_signature']='';bio._save(self.settings,token,record)
        with self.assertRaises(bio.ValidationError):self.import_plan(token)

    def test_visible7_proposal_sheet_is_secure_compatible_extension(self):
        self.execute("update person set birthday='1945';update rewards set number=18")
        original=bio.export_workbook(self.settings,'',20);wb=load_workbook(BytesIO(original))
        wb.active.cell(2,5,'Проверенная биография.');wb.active.cell(2,6,'https://warheroes.ru/hero/hero.asp?Hero_id=1');wb.active.cell(2,7,'Требует проверки')
        sheet=wb.create_sheet(bio.PROPOSALS_SHEET)
        sheet.append(['ID кавалера','ФИО в источнике','Год в источнике','Награда в источнике','Номер в источнике','Источник','Подтверждённые факты'])
        sheet.append([11095,'Ермакова (Зайцева) Мария Владимировна','1930','Орден Трудовой Славы I степени','18','https://warheroes.ru/hero/hero.asp?Hero_id=1','ФИО, I степень №18 и год1930 указаны в источнике.'])
        output=BytesIO();wb.save(output);wb.close()
        # Editing tools may reserialize blank metadata cells. The fixture uses
        # the untouched original hidden XML, just as an immutable-source merge.
        merged=BytesIO()
        with ZipFile(BytesIO(original)) as original_zip,ZipFile(BytesIO(output.getvalue())) as edited,ZipFile(merged,'w',ZIP_DEFLATED) as target:
            for part in edited.infolist():target.writestr(part,original_zip.read(part.filename) if part.filename=='xl/worksheets/sheet2.xml' else edited.read(part.filename))
        content=merged.getvalue()
        with ZipFile(BytesIO(original)) as a,ZipFile(BytesIO(content)) as b:self.assertEqual(a.read('xl/worksheets/sheet2.xml'),b.read('xl/worksheets/sheet2.xml'))
        token,_=bio.preview(self.settings,content);op,_=self.import_plan(token);flow.commit_import(self.settings,op)
        self.assertEqual(self.query('select birthday,biography from person'),[('1930','Проверенная биография.')])
        parsed=bio.read_rows(content)[0]['values'];self.assertEqual(parsed['Год рождения'],'1945')

    def test_proposal_duplicates_unknown_id_and_formulas_are_rejected(self):
        for value in ['=1+1',999999]:
            wb=load_workbook(BytesIO(bio.export_workbook(self.settings,'',20)))
            sheet=wb.create_sheet(bio.PROPOSALS_SHEET)
            sheet.append(['ID кавалера','ФИО в источнике','Год в источнике','Награда в источнике','Номер в источнике','Источник','Подтверждённые факты'])
            sheet.append([value,'Тестовый','1930','I','18','https://warheroes.ru/hero/hero.asp?Hero_id=1','Проверка'])
            out=BytesIO();wb.save(out);wb.close()
            with self.assertRaises(bio.ValidationError):bio.preview(self.settings,out.getvalue())

    def test_http_requires_exact_confirmation_and_shows_confirmed_disabled(self):
        token,_=self.preview();base='/data-maintenance/enrichment/biographies'
        with TestClient(app) as client:
            html=client.get(f'{base}/preview/{token}').text
            self.assertNotIn('Строка файла',html)
            response=client.post(f'{base}/import/{token}/plan',data={'revision':flow.present(self.settings,token)['revision']})
            self.assertEqual(response.status_code,200);self.assertIn('1930',response.text)
            op,plan=self.import_plan(token)
            self.assertEqual(client.post(base+'/import/commit',data={'operation':op}).status_code,409)
            self.assertEqual(client.post(base+'/import/commit',data={'operation':op,'confirm':'yes'},headers={'origin':'https://foreign.example'}).status_code,403)
            self.assertEqual(client.post(base+'/import/commit',data={'operation':op,'confirm':'yes'}).status_code,200)
            html=client.get(f'{base}/state/{token}').text
            self.assertIn('disabled aria-disabled="true" class="confirmed-button">Подтверждено',html)
            self.assertNotIn('data-manual-person=',html)
            history=client.get(f'{base}/review/{token}/2/history').text
            self.assertIn('Пользователь приложения',history);self.assertIn('Импорт',history)
            self.assertNotIn('manual_committed',history);self.assertNotIn('snapshot',history)

    def test_readonly_guard_and_history_no_empty_to_empty(self):
        token,_=self.preview();op,_=self.import_plan(token)
        with self.assertRaises(WriteBlockedError):flow.commit_import(self.settings.model_copy(update={'read_only':True}),op)
        flow.commit_import(self.settings,op)
        actions=flow.history_view(manual.durable_history(self.settings,11095))
        self.assertEqual({f['name'] for f in actions[0]['fields']},{'biography','birthday'})
        self.assertRegex(actions[0]['when'],r'\d{2}\.\d{2}\.\d{4}')
        self.assertEqual(self.media.read_bytes(),b'untouched TEMP media')

    def test_late_batch_failure_rolls_back_earlier_cavalier(self):
        self.execute("update person set birthday='1945';update rewards set number=18;"
                     "insert into person(id,fio,birthday,id_rank) values(2,'Тестов Второй Иванович','1945',10);"
                     "insert into rewards(person_id,id_gos,id_catigory,id_sub_catigory,id_name,number) values(2,1,2,3,4,19)")
        entries=bio.read_rows(self.content());rows=[]
        for e in entries:
            v=e['values']
            if v['person_id']=='2':v.update(source_url='https://warheroes.ru/hero/hero.asp?Hero_id=2',notes='ФИО и I степень №19 подтверждены; родился1930.')
            rows.append([v.get(h,'') for h in bio.HEADERS])
        token,_=bio.preview(self.settings,bio._workbook(bio.HEADERS,rows));op,_=self.import_plan(token)
        self.execute('''create trigger fail_second before insert on logs when instr(new.comment,'"person_id":11095')>0 begin select raise(abort,'second TEMP row');end;''')
        with self.assertRaises(sqlite3.IntegrityError):flow.commit_import(self.settings,op)
        self.assertEqual(self.query('select distinct birthday,biography from person'),[('1945',None)])
        self.assertEqual(self.query('select count(*) from logs'),[(0,)])

    def test_postponed_rows_are_excluded_from_explicit_bulk_plan(self):
        token,_=self.preview()
        with self.assertRaises(bio.ValidationError):flow.prepare_import(self.settings,token,flow.present(self.settings,token)['revision'],['2'])

    def test_obsolete_individual_1945_choice_is_not_silently_overridden(self):
        token,_=self.preview()
        with self.assertRaisesRegex(bio.ValidationError,'не требует ручного выбора'):
            self.choose(token,birthday_choice='current')

    def test_surname_or_initials_alone_are_not_confirmed_identity(self):
        for fio in ['Иванов','Иванов И.И.','Ермакова (Зайцева)']:
            self.execute("update person set fio='"+fio+"'")
            token,_=self.preview(fio=fio)
            detail=choices.detail(self.settings,token,2)
            self.assertFalse(detail['automatic']);self.assertEqual(detail['state'],'blocked')

    def test_sqlite_date_affinity_records_actual_year_and_does_not_create_false_stale(self):
        with closing(sqlite3.connect(self.db)) as c:dump='\n'.join(c.iterdump())
        self.assertIn('birthday text',dump)
        other=self.root/'date-affinity.sqlite'
        with closing(sqlite3.connect(other)) as c:c.executescript(dump.replace('birthday text','birthday DATE'))
        self.db=other;self.settings=self.settings.model_copy(update={'rewards_db_path':other})
        token,_=self.preview(current='1945-05-09')
        op,_=self.import_plan(token);flow.commit_import(self.settings,op)
        self.assertEqual(self.query('select birthday,typeof(birthday) from person'),[(1930,'integer')])
        self.assertTrue(choices.detail(self.settings,token,2)['confirmed'])
        action=manual.durable_history(self.settings,11095)[0]
        self.assertEqual(action['fields']['birthday']['after'],1930)

    def test_ordinary_7column_import_and_old_endpoint_remain_biography_only(self):
        self.execute("update person set birthday='1918'")
        wb=load_workbook(BytesIO(bio.export_workbook(self.settings,'',20)))
        wb.active.cell(2,5,'Обычная проверенная биография.');wb.active.cell(2,6,'https://example.org/source');wb.active.cell(2,7,'Найдено')
        out=BytesIO();wb.save(out);wb.close()
        token,record=bio.preview(self.settings,out.getvalue())
        self.assertEqual(record['rows'][0]['kind'],bio.READY)
        self.assertEqual(set(choices.detail(self.settings,token,2)['automatic']),{'biography'})
        with TestClient(app) as client:
            response=client.post('/data-maintenance/enrichment/biographies/commit',data={'token':token,'confirm':'yes'})
            self.assertEqual(response.status_code,200)
            self.assertIn('class="confirmed-button">Подтверждено',response.text)
        self.assertEqual(self.query('select birthday,biography from person'),[('1918','Обычная проверенная биография.')])


if __name__=='__main__':unittest.main()
