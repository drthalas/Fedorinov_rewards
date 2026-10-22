from concurrent.futures import ThreadPoolExecutor
from contextlib import closing
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
import hashlib
import os
import sqlite3
import unittest

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app.config import Settings
from backend.app.main import app
from backend.app.services import biography_enrichment as bio
from backend.app.services.write_guard import WriteBlockedError
from tests.test_ale475_person_import import create_database, xlsx


class BiographyEnrichmentTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.db = self.root / 'test.sqlite'
        create_database(self.db)
        self.settings = Settings(rewards_data_dir=self.root, rewards_db_path=self.db)
        self.execute("update person set birthday='1918'; update rewards set id_gos=1,id_catigory=2,id_sub_catigory=3;")
        self.env = patch.dict(os.environ, {'REWARDS_DATA_DIR':str(self.root), 'REWARDS_DB_PATH':str(self.db),
            'WRITE_MODE':'true', 'READ_ONLY':'false', 'UPDATE_CHECK_ENABLED':'false'})
        self.env.start()
        self.addCleanup(self.env.stop)
        # The module's generated state also remains TEMP and is removed by tests.
        self.state = patch.object(bio, '_state', return_value=self.root)
        self.state.start()
        self.addCleanup(self.state.stop)
        (self.root/'photo.jpg').write_bytes(b'temp media sentinel')

    def execute(self, sql):
        with closing(sqlite3.connect(self.db)) as connection:
            connection.executescript(sql)

    def query(self, sql):
        with closing(sqlite3.connect(self.db)) as connection:
            return connection.execute(sql).fetchall()

    def add(self, count=1):
        for person_id in range(2, count+2):
            self.execute(f"insert into person(id,fio,birthday,id_rank,comment,person_foto) values({person_id},'Тестов {person_id}','1918',10,'keep','photo.jpg');"
                         f"insert into rewards(person_id,id_gos,id_catigory,id_sub_catigory,id_name,number) values({person_id},1,2,3,4,{person_id});")

    def exported(self, limit=20):
        # Retain the original backend safety cases against the supported legacy
        # format. New contract cases below exercise the actual visible/hidden file.
        entries=bio.read_rows(self.human_exported(limit))
        return bio._workbook(bio.HEADERS, [[entry['values'].get(h,'') for h in bio.HEADERS] for entry in entries])

    def human_exported(self, limit=20):
        return bio.export_workbook(self.settings, '1/2/3/4', limit)

    def human_filled(self, edits=None, content=None):
        workbook=load_workbook(BytesIO(content or self.human_exported()))
        sheet=workbook.active
        headers={cell.value:cell.column for cell in sheet[1]}
        for number in range(2,sheet.max_row+1):
            for field,value in {'Биография':'Проверенная биография.','Источник':'https://warheroes.ru/hero/hero.asp?Hero_id=1','Статус':'Найдено'}.items():
                sheet.cell(number,headers[field],value)
        if edits: edits(sheet,headers,workbook)
        output=BytesIO();workbook.save(output);workbook.close()
        return output.getvalue()

    def filled(self, edits=None, content=None):
        workbook = load_workbook(BytesIO(content or self.exported()))
        sheet = workbook.active
        headers = {c.value:c.column for c in sheet[1]}
        for number in range(2, sheet.max_row+1):
            values = {name:sheet.cell(number,col).value for name,col in headers.items()}
            for name, value in {
                'biography':' Биография с переносом\nи фактами. ', 'source_url':'https://warheroes.ru/hero/hero.asp?Hero_id=1',
                'status':'Биография подтверждена','notes':'Подтверждённые факты источника',
                'ФИО в карточке':values['ФИО'],'Год в карточке':values['Год рождения'],
                'Звание в карточке':values['Звание'],'Награда в карточке':values['Награда'],
                'Идентификация':'ФИО, реальный год и награда совпали',
            }.items():
                sheet.cell(number,headers[name],value)
        if edits:
            edits(sheet, headers)
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        return stream.getvalue()

    def preview(self, **kwargs):
        return bio.preview(self.settings, self.filled(**kwargs))

    def test_unique_unicode_statistics_and_stable_read_only_export(self):
        self.add(3)
        self.execute("update person set biography=char(160)||char(8195) where id=2; update person set biography='Есть' where id=3;"
                     "insert into rewards(person_id,id_gos,id_catigory,id_sub_catigory,id_name,number) values(1,1,2,3,4,1000);")
        before = hashlib.sha256(self.db.read_bytes()).hexdigest()
        stats = bio.statistics(self.settings)
        self.assertEqual((stats['total'],stats['empty'],stats['filled']),(4,3,1))
        self.assertTrue(all(g['total']==4 and g['empty']==3 for g in stats['groups']))
        content = self.exported()
        rows = bio.read_rows(content)
        self.assertEqual([r['values']['person_id'] for r in rows],['1','2','4'])
        self.assertEqual(rows[0]['values']['Номер(а) награды'],'99 | 1000')
        self.assertEqual(hashlib.sha256(self.db.read_bytes()).hexdigest(),before)
        self.assertEqual(rows,bio.read_rows(self.exported()))

    def test_confirmed_commit_preserves_all_other_fields_rewards_and_media(self):
        self.add(1)
        before = self.query('select id,fio,birthday,id_rank,comment,person_foto from person'), self.query('select * from rewards')
        token, result = self.preview()
        self.assertEqual(result['ready'],2)
        self.assertEqual(self.query('select biography from person'),[(None,),(None,)])
        result = bio.commit(self.settings, token)
        self.assertEqual((result['saved'],result['state']),(2,'committed'))
        self.assertEqual(before,(self.query('select id,fio,birthday,id_rank,comment,person_foto from person'),self.query('select * from rewards')))
        self.assertEqual(self.query('select biography from person')[0][0],' Биография с переносом\nи фактами. ')
        self.assertEqual((self.root/'photo.jpg').read_bytes(),b'temp media sentinel')
        self.assertEqual(bio.commit(self.settings,token),result)

    def test_repeat_file_is_safe_and_preserves_filled_text(self):
        content = self.filled()
        token, _ = bio.preview(self.settings,content)
        bio.commit(self.settings,token)
        token2, result = bio.preview(self.settings,content)
        self.assertEqual(result['rows'][0]['kind'],bio.SKIP)
        self.assertEqual(bio.commit(self.settings,token2)['saved'],0)

    def test_duplicate_ids_block_every_occurrence(self):
        self.add()
        def edit(s,h): s.cell(3,h['person_id'],'1')
        token,result = self.preview(edits=edit)
        self.assertEqual(result['ready'],0)
        self.assertTrue(all(r['kind']==bio.CONFLICT for r in result['rows']))
        self.assertEqual(bio.commit(self.settings,token)['saved'],0)

    def test_identity_and_provenance_conflicts(self):
        for field,value in [('person_id','999'),('ФИО','Другой'),('Год рождения','1900'),('Награда','Другая'),
            ('Иерархия ID','1 / 2 / 3 / 5'),('Номер(а) награды','11'),('Звание','Капитан'),
            ('ФИО в карточке','Другой'),('Год в карточке','1900'),('Награда в карточке','Другая'),
            ('Звание в карточке','Капитан'),('Идентификация',''),('notes',''),('source_url','javascript:alert(1)'),
            ('status','Требует проверки'),('snapshot','{}'),('snapshot_signature','0'*64)]:
            with self.subTest(field=field):
                _,result = self.preview(edits=lambda s,h:s.cell(2,h[field],value))
                self.assertEqual(result['ready'],0)

    def test_default_1945_missing_year_empty_and_whitespace_biographies(self):
        for birthday in ['1945','','неизвестно']:
            self.execute(f"update person set birthday='{birthday}'")
            _,result = self.preview()
            self.assertEqual(result['ready'],0)
            self.assertIn('Год рождения',result['rows'][0]['reason'])
        self.execute("update person set birthday='1918'")
        for value in ['', '\u00a0\u2003']:
            _,result = self.preview(edits=lambda s,h:s.cell(2,h['biography'],value))
            self.assertEqual(result['rows'][0]['kind'],bio.SKIP)

    def test_unicode_cap_counts_spaces_and_linebreaks_without_truncation(self):
        for length,expected in [(2500,1),(2501,0)]:
            content = self.filled(edits=lambda s,h:s.cell(2,h['biography'],' '+ '😀'*(length-2)+'\n'))
            _,result = bio.preview(self.settings,content)
            self.assertEqual(result['ready'],expected)
            self.assertEqual(len(result['rows'][0]['biography']),length)

    def test_formula_and_error_cells_including_unmapped_columns_are_blocked(self):
        for field,value in [('biography','=1+1'),('notes','#VALUE!'),('snapshot','=1+1')]:
            _,result = self.preview(edits=lambda s,h:s.cell(2,h[field],value))
            self.assertEqual(result['rows'][0]['kind'],bio.ERROR)

    def test_malformed_files_limits_and_headers(self):
        for content in [b'',b'not zip',self.exported()[:100],xlsx([['1']],["person_id"])]:
            with self.assertRaises(bio.ValidationError): bio.preview(self.settings,content)
        self.add()
        content=self.filled()
        for name,value in [('MAX_BYTES',1),('MAX_UNPACKED',1),('MAX_ROWS',1),('MAX_COLUMNS',2)]:
            with patch.object(bio,name,value),self.assertRaises(bio.ValidationError): bio.preview(self.settings,content)

    def test_changed_identity_after_export_or_preview_is_not_written(self):
        self.add()
        content = self.filled()
        token, _ = bio.preview(self.settings,content)
        self.execute("update person set fio='Новый' where id=1")
        result = bio.commit(self.settings,token)
        self.assertEqual(result['saved'],1)
        self.assertEqual(result['rows'][0]['kind'],bio.CONFLICT)
        self.assertEqual(self.query('select biography from person where id=1'),[(None,)])

    def test_concurrently_filled_biography_and_deleted_person(self):
        self.add(2)
        token, _ = self.preview()
        self.execute("update person set biography='Ручной текст' where id=1; delete from rewards where person_id=2; delete from person where id=2;")
        result = bio.commit(self.settings,token)
        self.assertEqual(result['saved'],1)
        self.assertEqual(self.query('select biography from person where id=1'),[('Ручной текст',)])

    def test_write_guard_and_http_confirmation_origin_and_report(self):
        token, _ = self.preview()
        for update in [{'read_only':True},{'write_mode':False}]:
            with self.assertRaises(WriteBlockedError): bio.commit(self.settings.model_copy(update=update),token)
        with TestClient(app) as client:
            self.assertEqual(client.post('/data-maintenance/enrichment/biographies/commit',data={'token':token}).status_code,400)
            self.assertEqual(client.post('/data-maintenance/enrichment/biographies/commit',data={'token':token,'confirm':'yes'},headers={'origin':'https://other.example'}).status_code,403)
            report=client.get('/data-maintenance/enrichment/biographies/report/'+token)
            self.assertEqual(report.status_code,200)
            self.assertIn('warheroes.ru',report.text)
            self.assertIn('Подтверждённые факты',report.text)
            with patch.dict(os.environ,{'READ_ONLY':'true'}):
                self.assertEqual(client.post('/data-maintenance/enrichment/biographies/commit',data={'token':token,'confirm':'yes'}).status_code,403)
                self.assertEqual(client.get('/data-maintenance/enrichment/biographies/export?selection=1/2/3/4').status_code,200)

    def test_rollback_on_late_sqlite_failure(self):
        self.add()
        token,_ = self.preview()
        self.execute("create trigger fail_second before update of biography on person when new.id=2 begin select raise(ABORT,'synthetic failure'); end;")
        with self.assertRaises(sqlite3.IntegrityError): bio.commit(self.settings,token)
        self.assertEqual(self.query('select biography from person'),[(None,),(None,)])
        self.assertEqual(bio.load_preview(self.settings,token)['state'],'preview')

    def test_report_storage_failure_prevents_commit(self):
        token,_ = self.preview()
        with patch.object(bio,'_save',side_effect=OSError('disk full')),self.assertRaises(OSError): bio.commit(self.settings,token)
        self.assertEqual(self.query('select biography from person'),[(None,)])

    def test_post_commit_report_failure_is_reported_as_saved_not_rollback(self):
        token,_ = self.preview()
        real_save=bio._save
        def save(settings, token, record):
            if record['state']=='committed':
                raise OSError('synthetic report failure after commit')
            real_save(settings,token,record)
        with patch.object(bio,'_save',side_effect=save):
            result=bio.commit(self.settings,token)
        self.assertEqual(result['saved'],1)
        self.assertIn('окончательный статус',result['report_warning'])
        self.assertIsNotNone(self.query('select biography from person')[0][0])
        self.assertEqual(bio.load_preview(self.settings,token)['state'],'commit_pending')
        with self.assertRaises(bio.ValidationError): bio.commit(self.settings,token)

    def test_ignored_update_rolls_back_batch(self):
        self.add()
        token,_ = self.preview()
        self.execute('create trigger ignore_second before update of biography on person when new.id=2 begin select raise(IGNORE); end;')
        with self.assertRaises(bio.ValidationError): bio.commit(self.settings,token)
        self.assertEqual(self.query('select biography from person'),[(None,),(None,)])

    def test_reordered_columns_and_identity_rows_still_match_by_id(self):
        self.add()
        rows=bio.read_rows(self.filled())
        headers=list(reversed(bio.HEADERS))
        content=xlsx([[entry['values'][h] for h in headers] for entry in reversed(rows)],headers)
        _,result=bio.preview(self.settings,content)
        self.assertEqual(result['ready'],2)

    def test_exported_database_strings_never_become_formulas(self):
        self.execute("update person set fio='=1+1'")
        rows=bio.read_rows(self.exported())
        self.assertFalse(rows[0]['invalid'])
        self.assertEqual(rows[0]['values']['ФИО'],'=1+1')

    def test_incomplete_birth_or_bad_url_is_a_row_conflict_not_server_error(self):
        for value in ['https://[broken','ftp://example.org','https://user:pass@example.org']:
            _,result=self.preview(edits=lambda s,h:s.cell(2,h['source_url'],value))
            self.assertEqual(result['ready'],0)

    def test_concurrent_confirmations_serialize_and_are_idempotent(self):
        self.add(10)
        content = self.filled()
        tokens=[bio.preview(self.settings,content)[0] for _ in range(2)]
        with ThreadPoolExecutor(max_workers=2) as pool:
            results=list(pool.map(lambda t:bio.commit(self.settings,t),tokens))
        self.assertEqual(sorted(r['saved'] for r in results),[0,11])

    def test_signed_snapshots_are_database_bound_and_tamper_detected(self):
        content=self.filled()
        (self.root/'signing.key').write_bytes(b'different-key')
        _,result=bio.preview(self.settings,content)
        self.assertEqual(result['ready'],0)

    def test_ale491_legacy_mapping_19_ready_one_conflict_requires_original_export(self):
        self.add(19)
        original=self.exported()
        filled=bio.read_rows(self.filled(content=original))
        headers=['sample_order','person_id','ФИО','Год рождения','Звание','Награда','Иерархия ID',
            'Номер(а) награды','Сохранённые URL','Текущая биография','Примечание к году','Биография',
            'WarHeroes URL','Статус','ФИО в карточке','Год в карточке','Звание в карточке',
            'Сверка звания','Идентификация','Подтверждённые факты']
        rows=[]
        for i,entry in enumerate(filled):
            v=entry['values']
            v.update({'sample_order':i+1,'Биография':v['biography'],'WarHeroes URL':v['source_url'],
                'Статус':v['status'],'Подтверждённые факты':v['notes'],
                'Сверка звания':'Звание SQLite отличается от вводной части карточки; сверка показана отдельно. Идентификация подтверждается ФИО, годом и наградой. SQLite: Красноармеец; WarHeroes (вводная часть): лейтенант.',
                'Звание в карточке':'лейтенант'})
            if i==19:
                v.update({'Биография':'','Год в карточке':'1900','Статус':'Требует проверки идентификации'})
            rows.append([v.get(h,'') for h in headers])
        content=xlsx(rows,headers)
        self.assertEqual(bio.preview(self.settings,content)[1]['ready'],0)
        token,result=bio.preview(self.settings,content,original)
        self.assertEqual(result['ready'],19)
        self.assertEqual(result['rows'][-1]['kind'],bio.CONFLICT)
        self.assertEqual(bio.commit(self.settings,token)['saved'],19)

    def test_newly_eligible_rows_after_preview_stay_excluded(self):
        self.add()
        content=self.filled()
        self.execute("update person set biography='Ручная' where id=2")
        token,result=bio.preview(self.settings,content)
        self.assertEqual(result['ready'],1)
        self.execute("update person set biography=NULL where id=2")
        self.assertEqual(bio.commit(self.settings,token)['saved'],1)
        self.assertEqual(self.query('select biography from person where id=2'),[(None,)])

    def test_expired_and_invalid_preview_tokens(self):
        token,_=self.preview()
        with patch.object(bio.time,'time',return_value=10**12),self.assertRaises(bio.ValidationError):
            bio.load_preview(self.settings,token)
        for token in ['../signing.key','bad','f'*64]:
            with self.assertRaises(bio.ValidationError): bio.load_preview(self.settings,token)

    def test_http_navigation_upload_report_and_commit(self):
        with TestClient(app) as client:
            for path in ['/data-maintenance','/data-maintenance/enrichment','/data-maintenance/enrichment/biographies']:
                response=client.get(path)
                self.assertEqual(response.status_code,200)
                self.assertIn('Обслуживание данных',response.text)
            content=self.filled()
            response=client.post('/data-maintenance/enrichment/biographies/preview',files={'file':('filled.xlsx',content)})
            self.assertEqual(response.status_code,200)
            self.assertIn('Готово к сохранению',response.text)
            self.assertIn('Подтверждаю, что источники относятся к этим кавалерам',response.text)
            self.assertEqual(client.post('/data-maintenance/enrichment/biographies/preview',files={'file':('bad.xls',content)}).status_code,400)

    def test_human_visible_columns_hidden_metadata_and_roundtrip(self):
        self.add()
        self.execute('insert into rewards(person_id,id_gos,id_catigory,id_sub_catigory,id_name,number) values(1,1,2,3,4,222)')
        content=self.human_exported()
        workbook=load_workbook(BytesIO(content))
        self.assertEqual([c.value for c in workbook.active[1]],bio.VISIBLE_HEADERS)
        self.assertEqual([s.title for s in workbook if s.sheet_state=='visible'],['Биографии'])
        self.assertEqual(workbook[bio.METADATA_SHEET].sheet_state,'hidden')
        self.assertEqual(len(workbook.active.data_validations.dataValidation),1)
        workbook.close()
        entries=bio.read_rows(content)
        self.assertEqual(entries[0]['values']['Номер(а) награды'],'99 | 222')
        self.assertEqual(entries[0]['values']['Иерархия ID'],'1 / 2 / 3 / 4')
        filled=self.human_filled(content=content)
        token,result=bio.preview(self.settings,filled)
        self.assertEqual(result['ready'],2)
        self.assertEqual(bio.commit(self.settings,token)['saved'],2)
        self.assertEqual(bio.preview(self.settings,filled)[1]['ready'],0)

    def test_human_identity_fields_and_hidden_snapshot_cannot_be_changed(self):
        for field,value in [('ID кавалера','999'),('ФИО','Другой'),('Год рождения','1900'),('Награда','Другая')]:
            content=self.human_filled(edits=lambda s,h,w:s.cell(2,h[field],value))
            _,result=bio.preview(self.settings,content)
            self.assertEqual(result['ready'],0)
        for field,value in [('snapshot','{}'),('snapshot_signature','0'*64),('Звание','Капитан'),('Иерархия ID','1 / 2 / 3 / 5')]:
            def edit(s,h,w):
                hidden=w[bio.METADATA_SHEET];cols={c.value:c.column for c in hidden[1]}
                hidden.cell(2,cols[field],value)
            _,result=bio.preview(self.settings,self.human_filled(edits=edit))
            self.assertEqual(result['ready'],0)

    def test_human_metadata_missing_duplicate_formula_and_legacy_signature_fail_closed(self):
        for mode in ['missing','duplicate','formula']:
            def edit(s,h,w):
                hidden=w[bio.METADATA_SHEET]
                if mode=='missing': del w[bio.METADATA_SHEET]
                elif mode=='duplicate': hidden.append([c.value for c in hidden[2]])
                else: hidden.cell(2,2,'=1+1')
            with self.assertRaises(bio.ValidationError):bio.preview(self.settings,self.human_filled(edits=edit))
        self.assertEqual(self.query('select biography from person'),[(None,)])
        def legacy_signature(s,h,w):
            import json
            hidden=w[bio.METADATA_SHEET];cols={cell.value:cell.column for cell in hidden[1]}
            old=json.loads(hidden.cell(2,cols['snapshot']).value)
            old['version']=1;old.pop('contract')
            serialized=bio.canonical(old)
            hidden.cell(2,cols['snapshot'],serialized)
            hidden.cell(2,cols['snapshot_signature'],bio._sign(self.settings,serialized))
        _,result=bio.preview(self.settings,self.human_filled(edits=legacy_signature))
        self.assertEqual(result['ready'],0)
        self.assertIn('другому формату',result['rows'][0]['reason'])

    def test_human_status_semantics_and_three_categories(self):
        self.add(3)
        def edit(s,h,w):
            s.cell(3,h['Статус'],'Требует проверки')
            s.cell(4,h['Статус'],'Неизвестное значение')
            s.cell(5,h['Биография'],'')
            s.cell(5,h['Статус'],'')
        token,result=bio.preview(self.settings,self.human_filled(edits=edit))
        groups=bio.review_groups(result)
        self.assertEqual([len(rows) for rows in groups.values()],[1,2,1])
        self.assertIn('Неизвестный статус',groups[bio.REVIEW_LABEL][1]['reason'])
        self.assertEqual(bio.commit(self.settings,token)['saved'],1)
        self.assertEqual(self.query('select count(*) from person where biography is not null'),[(1,)])

    def test_human_default_year_filled_missing_and_duplicate_id(self):
        self.add(3)
        self.execute("update person set birthday='1945' where id=2")
        content=self.human_filled(edits=lambda s,h,w:s.cell(5,h['ID кавалера'],'999'))
        self.execute("update person set biography='Существующая' where id=3")
        _,result=bio.preview(self.settings,content)
        self.assertEqual([r['kind'] for r in result['rows']],[bio.READY,bio.CONFLICT,bio.SKIP,bio.SKIP])
        content=self.human_filled(edits=lambda s,h,w:s.append([c.value for c in s[2]]))
        _,result=bio.preview(self.settings,content)
        self.assertIn('повторяется',result['rows'][0]['reason'])

    def test_human_concurrency_rollback_and_only_biography_updates(self):
        self.add()
        content=self.human_filled()
        before=self.query('select id,fio,birthday,id_rank,comment,person_foto from person'),self.query('select * from rewards')
        token,_=bio.preview(self.settings,content)
        self.execute("create trigger fail_second before update of biography on person when new.id=2 begin select raise(ABORT,'test'); end;")
        with self.assertRaises(sqlite3.IntegrityError):bio.commit(self.settings,token)
        self.assertEqual(self.query('select biography from person'),[(None,),(None,)])
        self.execute("drop trigger fail_second; update person set biography='Ручная запись' where id=2")
        self.assertEqual(bio.commit(self.settings,token)['saved'],1)
        self.assertEqual(before,(self.query('select id,fio,birthday,id_rank,comment,person_foto from person'),self.query('select * from rewards')))
        self.assertEqual(self.query('select biography from person where id=2'),[('Ручная запись',)])

    def test_human_legacy_selection_mismatch_has_instructions(self):
        self.add()
        old=bio.read_rows(self.filled())[1]['values']
        old.pop('snapshot');old.pop('snapshot_signature')
        content=xlsx([[old.get(h,'') for h in bio.HEADERS]],bio.HEADERS)
        original=self.human_exported(limit=1)
        _,result=bio.preview(self.settings,content,original)
        self.assertEqual(result['ready'],0)
        self.assertIn('исходный Excel с теми же кавалерами',result['rows'][0]['reason'])

    def test_landing_disabled_placeholder_and_user_facing_copy(self):
        with TestClient(app) as client:
            landing=client.get('/data-maintenance').text
            self.assertIn('Оптимизация изображений',landing)
            self.assertIn('В разработке',landing)
            self.assertIn('aria-disabled="true"',landing)
            response=client.post('/data-maintenance/enrichment/biographies/preview',files={'file':('new.xlsx',self.human_filled())})
            for label in ['1. Скачать Excel','2. Загрузить заполненный Excel','3. Проверить и сохранить',bio.READY_LABEL,bio.REVIEW_LABEL,bio.SKIP_LABEL]:
                self.assertIn(label,response.text)
            for internal in ['пилот','ALE-491','snapshot','source_url','notes','research']:
                self.assertNotIn(internal,response.text)

    def test_human_excel_report_contains_categories_reasons_and_source(self):
        token,result=bio.preview(self.settings,self.human_filled())
        with TestClient(app) as client:
            response=client.get('/data-maintenance/enrichment/biographies/report/'+token+'/excel')
            self.assertEqual(response.status_code,200)
        workbook=load_workbook(BytesIO(response.content))
        row=[cell.value for cell in workbook.active[2]]
        self.assertEqual(row[3],bio.READY_LABEL)
        self.assertIn('warheroes.ru',row[6])
        self.assertTrue(row[4])
        workbook.close()


if __name__ == '__main__':
    unittest.main()
