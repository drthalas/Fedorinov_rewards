"""Current two-bucket contract; signed TEMP fixtures only."""
from contextlib import closing
from io import BytesIO
import json,sqlite3,unittest
from unittest.mock import patch
from tests import test_ale513_review_choices as fixture
from backend.app.services import biography_enrichment as b,manual_review_choices as c,enrichment_workflow as f
from backend.app.services import manual_import_review as manual
from backend.app.main import app
from fastapi.testclient import TestClient

class AutoImportTests(unittest.TestCase):
 setUp=fixture.ImportWorkflowTests.setUp
 execute=fixture.ImportWorkflowTests.execute
 query=fixture.ImportWorkflowTests.query
 content=fixture.ImportWorkflowTests.content
 preview=fixture.ImportWorkflowTests.preview
 def commit(self,t):
  v=f.workspace(self.settings,t);op,p=f.prepare_workspace(self.settings,t,v['revision'],[]);return f.commit_import(self.settings,op),op
 def test_non1945_and_1945_source_years_auto_no_fio_rewrite(self):
  for current,new in [('1932','1931'),('1945','1930')]:
   with self.subTest(current=current):
    self.execute('delete from logs;update person set biography=null;')
    t,_=self.preview(current=current,year=new);d=c.detail(self.settings,t,2)
    self.assertEqual(d['fields'],{});self.assertEqual(set(d['automatic']),{'biography','birthday'})
    before=self.db.read_bytes();v=f.workspace(self.settings,t);op,p=f.prepare_workspace(self.settings,t,v['revision'],[]);self.assertEqual(self.db.read_bytes(),before)
    r=f.commit_import(self.settings,op);self.assertEqual(f.commit_import(self.settings,op),r)
    self.assertEqual(self.query('select birthday,fio from person'),[(new,'Ермакова (Зайцева) Мария Владимировна')])
    self.assertEqual(f.workspace(self.settings,t)['counts'],{'ready':0,'insufficient':0})
 def test_alias_is_automatic_only_with_exact_serial_and_year(self):
  t,_=self.preview(current='1954',year='1954',fio='Ермакова Мария Владимировна',identity='Личность не подтверждена')
  self.assertEqual(set(c.detail(self.settings,t,2)['automatic']),{'biography'})
  self.commit(t);self.assertEqual(self.query('select fio from person'),[('Ермакова (Зайцева) Мария Владимировна',)])
 def test_true1945_existing_bio_and_no_changes_are_not_missing_evidence(self):
  t,_=self.preview(current='1945-05-09',year='1945');self.commit(t)
  self.assertEqual(self.query('select birthday from person'),[('1945-05-09',)])
  self.execute('delete from logs')
  t,_=b.preview(self.settings,self.last_content);v=f.workspace(self.settings,t)
  self.assertEqual(v['counts'],{'ready':0,'insufficient':0});self.assertEqual(len(v['unchanged']),1)
  self.assertEqual(self.query('select biography from person'),[('Проверенная биография.',)])
 def test_missing_bio_and_wrong_identity_remain_insufficient(self):
  for kwargs in [{'biography':''},{'fio':'Совсем Другой Человек'},{'award':'Трудовая Слава II степени'}]:
   t,_=self.preview(current='1932',year='1932',**kwargs);d=c.detail(self.settings,t,2)
   self.assertFalse(d['automatic']);self.assertEqual(f.workspace(self.settings,t)['counts']['insufficient'],1)
  with self.assertRaises(b.ValidationError):f.stage(self.settings,t,2,{'identity_choice':'same'})
 def test_gss_summary_can_prove_bio_but_never_guesses_source_year(self):
  self.execute("update guide_lev_3 set name='Герой Советского Союза II' where id=4;update person set fio='Ворожейкин Арсений Васильевич',birthday='1912';update rewards set number=28")
  native=b.export_workbook(self.settings,'',20)
  from openpyxl import load_workbook
  w=load_workbook(BytesIO(native));s=w.active;s.cell(2,5,'Арсений Васильевич Ворожейкин (9 ноября 1914 — 23 мая 2001) — лётчик, дважды Герой Советского Союза. Награждён второй медалью «Золотая Звезда» № 28.');s.cell(2,6,'https://warheroes.ru/hero/hero.asp?Hero_id=1226');s.cell(2,7,'Требует проверки')
  o=BytesIO();w.save(o);w.close();t,_=b.preview(self.settings,o.getvalue());d=c.detail(self.settings,t,2)
  self.assertEqual(set(d['automatic']),{'biography'});self.commit(t);self.assertEqual(self.query('select birthday from person'),[('1912',)])
 def test_year_caveat_preserves_year_but_allows_bio(self):
  t,_=self.preview(current='1912',year='1914');r=b.load_preview(self.settings,t);r['entries'][0]['values']['Год подтверждён']='нет';b._save(self.settings,t,r)
  self.assertEqual(set(c.detail(self.settings,t,2)['automatic']),{'biography'})
 def test_hmac_stale_atomic_guards_and_no_source_fio_api(self):
  t,_=self.preview();v=f.workspace(self.settings,t);op,_=f.prepare_workspace(self.settings,t,v['revision'],[])
  self.execute("update person set biography='Чужой текст'")
  with self.assertRaises(b.ValidationError):f.commit_import(self.settings,op)
  self.execute('update person set biography=null')
  t,_=self.preview();record=b.load_preview(self.settings,t);record['entries'][0]['values']['snapshot_signature']='0'*64;b._save(self.settings,t,record)
  self.assertTrue(f.workspace(self.settings,t)['fatal'])
  with self.assertRaises(b.ValidationError):f.prepare_workspace(self.settings,t,f.workspace(self.settings,t)['revision'],[])
  with self.assertRaises(b.ValidationError):f.commit_review(self.settings,op)
 def test_fio_correction_invalidates_original_signed_snapshot(self):
  from backend.app.services import fio_corrections
  self.execute("update person set fio='Cидоров Петр',birthday='1918'")
  raw=b.export_workbook(self.settings,'',20);token,_=b.preview(self.settings,raw)
  settings=self.settings.model_copy(update={'update_backup_dir':self.root/'correction-backups'})
  audit=fio_corrections.audit(settings);fio_corrections.apply(settings,audit['audit_token'],'9'*64)
  self.assertEqual(self.query('select fio from person'),[('Сидоров Петр',)])
  self.assertTrue(f.workspace(self.settings,token)['fatal'])
  with self.assertRaises(b.ValidationError):f.prepare_workspace(self.settings,token,f.workspace(self.settings,token)['revision'],[])

 def test_material_identity_contradiction_cannot_be_overridden(self):
  t,_=self.preview(identity='Другой человек')
  self.assertFalse(c.detail(self.settings,t,2)['automatic'])
  self.assertEqual(f.workspace(self.settings,t)['counts']['insufficient'],1)

 def test_two_tabs_and_no_manual_actions(self):
  t,_=self.preview(current='1932',year='1931')
  with TestClient(app) as client:
   html=client.get('/data-maintenance/enrichment/biographies/preview/'+t).text
   self.assertEqual(html.count('role="tab"'),2);self.assertNotIn('Требует решения',html);self.assertNotIn('Подтвердить данные',html)
   self.assertIn('Исправление данных',client.get('/data-maintenance').text)

if __name__=='__main__':unittest.main()
