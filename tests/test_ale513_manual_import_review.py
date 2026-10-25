"""Risk-based tests. All writes, fixtures and audit live in TemporaryDirectory."""
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
import json
import os
import sqlite3
import unittest

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app.config import Settings
from backend.app.main import app
from backend.app.services import biography_enrichment as bio, manual_import_review as manual, manual_review_choices as choices
from backend.app.services.write_guard import WriteBlockedError
from tests.test_ale475_person_import import create_database


class ManualReviewTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory(prefix='ale513-')
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.db = self.root/'test.sqlite'
        create_database(self.db)
        self.settings = Settings(rewards_data_dir=self.root, rewards_db_path=self.db)
        self.execute("create table logs(datetime DATETIME default current_timestamp, action VARCHAR, comment DATE);"
                     "update rewards set id_gos=1,id_catigory=2,id_sub_catigory=3;"
                     "update guide_lev_3 set name='Орден Трудовой Славы I степени' where id=4;"
                     "update person set id=11095,fio='Ермакова (Зайцева) Мария Владимировна',birthday='1918';"
                     "update rewards set person_id=11095,number=725;")
        self.state = patch.object(bio, '_state', return_value=self.root)
        self.state.start(); self.addCleanup(self.state.stop)
        self.env = patch.dict(os.environ, {'REWARDS_DATA_DIR': str(self.root), 'REWARDS_DB_PATH': str(self.db),
                              'WRITE_MODE':'true', 'READ_ONLY':'false', 'UPDATE_CHECK_ENABLED':'false'})
        self.env.start(); self.addCleanup(self.env.stop)
        self.media = self.root/'sentinel.jpg'; self.media.write_bytes(b'untouched TEMP media')

    def execute(self, sql):
        with closing(sqlite3.connect(self.db)) as c:
            c.executescript(sql)

    def query(self, sql):
        with closing(sqlite3.connect(self.db)) as c:
            return c.execute(sql).fetchall()

    def fixture(self, edits=None):
        workbook = load_workbook(BytesIO(bio.export_workbook(self.settings, '', 20)))
        s = workbook.active
        h = {c.value: c.column for c in s[1]}
        for field, value in {'Биография':'Тестовая проверенная биография.', 'Источник':'https://example.org/person/11095', 'Статус':'Требует проверки'}.items():
            s.cell(2,h[field],value)
        if edits: edits(workbook,s,h)
        output=BytesIO(); workbook.save(output); workbook.close()
        return output.getvalue()

    def preview(self, edits=None):
        return bio.preview(self.settings, self.fixture(edits))

    def submitted(self, **changes):
        values = {'actor':'Owner', 'identity_confirm':'yes'}
        for name in manual.FIELDS:
            values.update({name+'_decision':'pending', name+'_proposal':'', name+'_source':'https://example.org/person/11095',
                           name+'_evidence':'Проверены ФИО, год и орден I степени №725 по источнику.'})
        values.update(changes)
        return values

    def plan(self, token, **changes):
        return manual.prepare(self.settings, token, 2, self.submitted(**changes))

    def test_ermakova_biography_accept_alias_reject_year_pending_and_durable_audit(self):
        token, _ = self.preview()
        context=manual.review(self.settings,token,2)
        self.assertEqual(context['identity']['awards'][0]['number'],725)
        self.assertEqual(context['identity']['awards'][0]['award'],'Орден Трудовой Славы I степени')
        before=self.db.read_bytes()
        p,plan=self.plan(token, biography_decision='accept',biography_proposal='Отдельно принятая биография.',
                         fio_decision='reject',fio_proposal='Ермакова Мария Владимировна')
        self.assertEqual(self.db.read_bytes(),before)  # Decisions alone never write SQLite.
        event=manual.commit(self.settings,p)
        self.assertEqual(event['writes'],['biography'])
        self.assertEqual(self.query('select fio,birthday,biography from person'),
            [('Ермакова (Зайцева) Мария Владимировна','1918','Отдельно принятая биография.')])
        self.assertEqual(event,manual.commit(self.settings,p))
        self.assertEqual(len(self.query('select * from logs')),1)
        self.assertEqual(event['fields']['fio']['decision'],'reject')
        self.assertEqual(event['fields']['birthday']['decision'],'pending')
        self.assertEqual(manual.reviewed_history(self.settings,token,2)[0],event)
        self.assertIn('ФИО: Отклонено',manual.summaries(self.settings,event['file_sha256'])[2])
        self.assertEqual(self.media.read_bytes(),b'untouched TEMP media')

    def test_year_corrections_and_true_1945(self):
        for fio,number,old,new in [('Кучкаров Тест Тестович',18,'1945','1938'),
                                   ('Солнцев Тест Тестович',502,'1932','1931'),
                                   ('Подтверждённый Тест Тестович',44,'1945','1945')]:
            with self.subTest(fio=fio):
                self.execute(f"update person set fio='{fio}',birthday='{old}',biography=null; update rewards set number={number};")
                token,result=self.preview()
                self.assertEqual(result['ready'],0)
                p,_=self.plan(token,birthday_decision='accept',birthday_proposal=new,biography_decision='reject',biography_proposal='Не принимать')
                event=manual.commit(self.settings,p)
                self.assertEqual(self.query('select birthday,biography from person'),[(new,None)])
                self.assertEqual(event['writes'],[] if old==new else ['birthday'])

    def test_multiple_fields_accept_and_history_after_stale_export(self):
        token,_=self.preview()
        frozen=self.query('select id,id_rank,comment,person_foto from person'),self.query('select * from rewards')
        p,_=self.plan(token,fio_decision='accept',fio_proposal='Ермакова Мария Владимировна',
                      birthday_decision='accept',birthday_proposal='1938',biography_decision='accept',biography_proposal='Подтверждённый текст')
        event=manual.commit(self.settings,p)
        self.assertEqual(set(event['writes']),set(manual.FIELDS))
        self.assertEqual(len(manual.reviewed_history(self.settings,token,2)),1)
        self.assertEqual(manual.review(self.settings,token,2)['current']['birthday'],'1938')
        self.assertEqual(event,manual.commit(self.settings,p))
        self.assertEqual(frozen,(self.query('select id,id_rank,comment,person_foto from person'),self.query('select * from rewards')))

    def test_external_biography_change_is_not_silently_rebased(self):
        token,_=self.preview()
        p,_=self.plan(token,birthday_decision='accept',birthday_proposal='1938')
        self.execute("update person set biography='Существующий текст'")
        with self.assertRaisesRegex(bio.ValidationError,'вне этого импорта'):
            manual.review(self.settings,token,2)
        with self.assertRaises(bio.ValidationError):manual.commit(self.settings,p)
        self.assertEqual(self.query('select biography,birthday from person'),[('Существующий текст','1918')])

    def preview_after_empty_export(self):
        self.execute('update person set biography=null')
        token,result=self.preview()
        self.execute("update person set biography='Существующий текст'")
        return token,result

    def test_hmac_other_installation_stale_identity_and_frozen_visible_fields_fail_closed(self):
        for edit in [lambda w,s,h:s.cell(2,h['ФИО'],'Подменённое Имя'),
                     lambda w,s,h:w[bio.METADATA_SHEET].cell(2,bio.HEADERS.index('snapshot_signature')+1,'0'*64)]:
            token,_=self.preview(edit)
            with self.assertRaises(bio.ValidationError): manual.review(self.settings,token,2)
        token,_=self.preview()
        self.execute('update rewards set number=999')
        with self.assertRaisesRegex(bio.ValidationError,'исходные данные устарели'): self.plan(token)
        self.assertEqual(self.query('select count(*) from logs'),[(0,)])

    def test_invalid_fields_and_missing_identity_evidence(self):
        token,_=self.preview()
        cases=[{'birthday_decision':'accept','birthday_proposal':y} for y in ['1799','2000','1945.0','']]
        cases += [{'fio_decision':'accept','fio_proposal':n} for n in ['Иван','Имя 123','<script> Имя']]
        cases += [{'biography_decision':'accept','biography_proposal':'😀'*2501},
                  {'fio_decision':'accept','fio_proposal':'Ермакова Мария Владимировна','identity_confirm':''},
                  {'birthday_decision':'accept','birthday_proposal':'1938','birthday_source':'javascript:bad'},
                  {'birthday_decision':'accept','birthday_proposal':'1938','birthday_evidence':''},
                  {'actor':'AI'}]
        before=self.db.read_bytes()
        for changes in cases:
            with self.subTest(changes=changes):
                with self.assertRaises(bio.ValidationError): self.plan(token,**changes)
        self.assertEqual(self.db.read_bytes(),before)

    def test_foreign_database_signing_key_missing_person_and_concurrent_award_change(self):
        content=self.fixture()
        other=self.root/'other-install';other.mkdir()
        with patch.object(bio,'_state',return_value=other):
            token,_=bio.preview(self.settings,content)
            with self.assertRaisesRegex(bio.ValidationError,'служебные данные'): manual.review(self.settings,token,2)
        token,_=self.preview()
        p,_=self.plan(token,birthday_decision='accept',birthday_proposal='1938')
        self.execute('update rewards set number=726')
        with self.assertRaisesRegex(bio.ValidationError,'Награды и номера'): manual.commit(self.settings,p)
        self.assertEqual(self.query('select count(*) from logs'),[(0,)])
        self.execute('delete from rewards; delete from person')
        with self.assertRaisesRegex(bio.ValidationError,'исходные данные устарели'): manual.commit(self.settings,p)

    def test_durable_history_survives_temp_preview_removal(self):
        token,_=self.preview()
        p,_=self.plan(token)
        event=manual.commit(self.settings,p)
        (self.root/(token+'.json')).unlink()
        self.assertEqual(manual.durable_history(self.settings,11095),[event])
        self.assertEqual(manual.durable_history(self.settings,987),[])

    def test_atomic_rollback_when_audit_fails_and_idempotent_concurrent_retry(self):
        token,_=self.preview()
        p,_=self.plan(token,birthday_decision='accept',birthday_proposal='1938',biography_decision='accept',biography_proposal='Принятая биография')
        self.execute("create trigger fail_audit before insert on logs begin select raise(abort,'TEMP audit failure'); end;")
        with self.assertRaises(sqlite3.IntegrityError): manual.commit(self.settings,p)
        self.assertEqual(self.query('select birthday,biography from person'),[('1918',None)])
        self.assertEqual(self.query('select count(*) from logs'),[(0,)])
        self.execute('drop trigger fail_audit')
        with ThreadPoolExecutor(2) as pool:
            events=list(pool.map(lambda _:manual.commit(self.settings,p),range(2)))
        self.assertEqual(events[0],events[1])
        self.assertEqual(self.query('select count(*) from logs'),[(1,)])

    def test_conflicting_plans_readonly_plan_tamper_and_schema_block(self):
        token,_=self.preview()
        p,_=self.plan(token,biography_decision='accept',biography_proposal='Первое решение')
        q,_=self.plan(token,biography_decision='accept',biography_proposal='Второе решение')
        manual.commit(self.settings,p)
        with self.assertRaises(bio.ValidationError): manual.commit(self.settings,q)
        with self.assertRaises(WriteBlockedError): manual.commit(self.settings.model_copy(update={'read_only':True}),p)
        payload=json.loads((self.root/(q+'.json')).read_text()); payload['fields']['fio']['decision']='accept'
        bio._save(self.settings,q,payload)
        with self.assertRaisesRegex(bio.ValidationError,'повреждён'): manual.commit(self.settings,q)
        self.execute('drop table logs')
        with self.assertRaisesRegex(bio.ValidationError,'миграции'): manual.review(self.settings,token,2)

    def test_http_blocker_has_no_pseudo_choices_and_origin_is_checked(self):
        token,_=self.preview()
        base='/data-maintenance/enrichment/biographies'
        with TestClient(app) as client:
            response=client.get(f'{base}/review/{token}/2')
            self.assertEqual(response.status_code,200)
            self.assertIn('Что мешает импорту',response.text)
            self.assertNotIn('type="radio"',response.text)
            self.assertEqual(client.post(f'{base}/review/{token}/2/plan',data={}).status_code,400)
            p,_=self.plan(token)
            self.assertEqual(client.post(base+'/review/commit',data={'plan_token':p}).status_code,409)
            self.assertEqual(client.post(base+'/review/commit',data={'plan_token':p,'confirm':'yes'},headers={'origin':'https://elsewhere.example'}).status_code,403)
            self.assertEqual(client.get(f'{base}/preview/{token}').status_code,200)
            self.assertEqual(client.get(base+'/review-history?person_id=11095').status_code,200)


if __name__ == '__main__':
    unittest.main()
